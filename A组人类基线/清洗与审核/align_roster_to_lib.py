# -*- coding: utf-8 -*-
"""把人类树与清洗数据的 M 名册对齐到「史料总库.xlsx · 史料库」权威名单。

背景：
  原文_clean.json 的 materials 来自论文脚注提取 + LLM 判定（见总库「汇总统计」的
  「史料条数」列），条数多于总库「史料库」工作表实际收录数，差额是在建库时
  「找不到对应条目」而主动删除的（A 组共 57 条）。

为什么必须删：
  三档史料池（S/M/L）由总库生成，模型侧只能看到总库收录的史料。
  人类树若保留总库外的 M，这些 C→M 边模型永远无法复现，会系统性压低 ECI。

本脚本只做「删条目 + 同步 sources」的确定性后处理，不重跑 LLM：
  C 的 claim 由段落文本决定，与 M 名册无关，故重抽不会改变 claim。

用法: python align_roster_to_lib.py [--apply]
  不带 --apply 为预演（只打印影响）；带 --apply 才写盘（先自动备份 .bak_名册对齐）。
"""
import json
import shutil
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(r'C:/Users/23297/Downloads/FDUROP')
LIB = ROOT / '史料总库.xlsx'
BASE = ROOT / r'归档/数据集/数据清洗后论文'
RUN = 'A_human_自然段基线'
PIDS = ['A01', 'A02', 'A03', 'A04', 'A05', 'A06', 'A07', 'A08', 'A09', 'A10']


def load_lib_roster():
    """总库「史料库」→ {篇目: {M 编号}}。"""
    wb = openpyxl.load_workbook(LIB, data_only=True)
    ws = wb['史料库']
    out = defaultdict(set)
    for i in range(2, ws.max_row + 1):
        a, c = ws.cell(i, 1).value, ws.cell(i, 3).value
        if a is None or c is None:
            continue
        out[str(a).strip()].add(str(c).strip())
    return out


def main():
    apply = '--apply' in sys.argv
    lib = load_lib_roster()
    rows, tg, ta, to = [], 0, 0, 0

    for pid in PIDS:
        cpath = BASE / pid / '00_原文' / '原文_clean_自然段.json'
        tpath = BASE / pid / '03_实验输出' / 'Q_A_实验' / RUN / 'run_论证树.json'
        clean = json.loads(cpath.read_text(encoding='utf-8'))
        tree = json.loads(tpath.read_text(encoding='utf-8'))
        keep = lib.get(pid, set())
        gone = {m['id'] for m in clean['materials']} - keep
        if not gone:
            rows.append((pid, 0, 0, 0, len(clean['materials'])))
            continue

        enc = tree['encoding_table']

        def purge(ids):
            return [x for x in ids if x not in gone]

        aff = only = 0
        for c in enc['C_level']:
            s = set(c.get('sources') or [])
            if s & gone:
                aff += 1
                if s <= gone:
                    only += 1
            c['sources'] = sorted(purge(c['sources']), key=lambda x: (len(x), x))
            c['source_count'] = len(c['sources'])
        for lv in ('B_level', 'A_level'):
            for n in enc[lv]:
                n['sources'] = sorted(purge(n['sources']), key=lambda x: (len(x), x))
                n['source_count'] = len(n['sources'])

        # 名册本身
        clean['materials'] = [m for m in clean['materials'] if m['id'] not in gone]
        for p in clean['paragraphs']:
            p['material_refs'] = [r for r in p.get('material_refs') or [] if r not in gone]
        enc['M_level'] = [m for m in enc['M_level'] if m['id'] not in gone]

        # 硬约束②：A1.sources 必须等于完整名册
        new_roster = sorted({m['id'] for m in clean['materials']}, key=lambda x: (len(x), x))
        for a in enc['A_level']:
            a['sources'] = list(new_roster)
            a['source_count'] = len(new_roster)

        # statistics
        st = tree['statistics']
        used = set()
        for c in enc['C_level']:
            used |= set(c['sources'])
        st['total_unique_M'] = len(new_roster)
        st['C_nodes_without_sources'] = sum(1 for c in enc['C_level'] if not c['sources'])
        st['M'] = {'total': len(new_roster),
                   'used': len(used),
                   'unused': sorted(set(new_roster) - used, key=lambda x: (len(x), x))}
        st['source_count_by_A'] = {a['id']: a['source_count'] for a in enc['A_level']}
        st['source_count_by_B'] = {b['id']: b['source_count'] for b in enc['B_level']}
        tree['meta']['roster_aligned_to_lib'] = {
            'removed': sorted(gone), 'kept': len(new_roster),
            'note': 'M 名册对齐到史料总库「史料库」权威名单（总库外条目为建库时找不到而删除者）',
        }

        if apply:
            for p in (cpath, tpath):
                bak = p.with_suffix(p.suffix + '.bak_名册对齐')
                if not bak.exists():
                    shutil.copy2(p, bak)
            cpath.write_text(json.dumps(clean, ensure_ascii=False, indent=1), encoding='utf-8')
            tpath.write_text(json.dumps(tree, ensure_ascii=False, indent=1), encoding='utf-8')

        rows.append((pid, len(gone), only, aff, len(new_roster)))
        tg += len(gone); to += only; ta += aff

    print(f'{"篇":<5}{"删除M":>7}{"新无史料C":>10}{"受影响C":>9}{"新名册":>8}')
    print('-' * 42)
    for pid, g, o, a, n in rows:
        print(f'{pid:<5}{g:>7}{o:>10}{a:>9}{n:>8}')
    print('-' * 42)
    print(f'{"合计":<5}{tg:>7}{to:>10}{ta:>9}')
    print(f'\n模式：{"已写盘（原文件备份为 .bak_名册对齐）" if apply else "预演，未写盘"}')


if __name__ == '__main__':
    main()
